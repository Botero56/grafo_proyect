"""Genera de verdad el Excel (sin mockear generate_report) y
revisa el archivo resultante. Los demás tests de reportes solo cubren los
resúmenes JSON; sin esto, el camino que escribe el Excel -- tablas, KPIs,
gráficas -- no lo ejecutaba ningún test."""

from datetime import datetime, timedelta

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from app.firestore_db.firebase_schemas import Orders
from app.services.excel_report_service import dashboard_report
from app.services.excel_report_service.constants import DEFAULT_TZ
from app.services.excel_report_service.excel_report_service import ExcelReportService
from app.services.excel_report_service.models import ExcelReportSpec
from tests.conftest import seed_business, seed_conversation, seed_customer

# Fecha fija: los reportes filtran por día/mes, así que "hoy" no puede
# depender de cuándo corre el test.
REFERENCE = datetime(2026, 9, 15, 12, 30, tzinfo=DEFAULT_TZ)


def seed_report_data(fake_store):
    """Datos que ejercitan todas las ramas: dos productos, un cancelado,
    días de comparación (ayer, semana pasada, mes anterior) y un cliente
    recurrente."""
    business = seed_business(fake_store)
    camilo = seed_customer(fake_store, business, "573001111111", name="Camilo")
    ana = seed_customer(fake_store, business, "573002222222", name="Ana")
    seed_conversation(fake_store, business, camilo.customer_id, channel="whatsapp")
    seed_conversation(fake_store, business, ana.customer_id, channel="whatsapp", status="closed")

    def order(menu_id, qty, price, status, when, customer):
        o = Orders()
        o.business_id = business.bussiness_id
        o.customer_id = customer.customer_id
        o.menu_id = menu_id
        o.quantity = qty
        o.unit_price = price
        o.total = qty * price
        o.status = status
        o.created_at = when
        fake_store.seed(o)

    order("menu_del_dia", 2, 18500, "delivered", REFERENCE, camilo)
    order("otras_opciones", 1, 20500, "delivered", REFERENCE + timedelta(hours=1), ana)
    order("menu_del_dia", 1, 18500, "cancelled", REFERENCE + timedelta(hours=2), ana)
    order("menu_del_dia", 3, 18500, "delivered", REFERENCE - timedelta(days=1), camilo)
    order("otras_opciones", 2, 20500, "delivered", REFERENCE - timedelta(days=7), ana)
    order("menu_del_dia", 1, 18500, "delivered", REFERENCE - timedelta(days=40), camilo)
    return business


def generate(business, tmp_path, report_type, **spec_fields):
    spec = ExcelReportSpec(
        business_id=business.bussiness_id,
        report_type=report_type,
        output_path=tmp_path / f"{report_type}.xlsx",
        **spec_fields,
    )
    path = ExcelReportService(connect=False).generate_report(spec)
    return load_workbook(path)


def _column(ws, col: int) -> list:
    return [ws.cell(row=r, column=col).value for r in range(1, ws.max_row + 1)]


def test_every_report_type_is_the_template_with_the_orders(fake_store, tmp_path, dashboard_template):
    """Sea cual sea el report_type pedido, el Excel es la plantilla con los
    pedidos del rango pegados en la tabla `datos`."""
    business = seed_report_data(fake_store)
    day = REFERENCE.replace(hour=0, minute=0)

    for report_type in ("daily_sales", "customers", "conversations", "lunch_stats"):
        wb = generate(business, tmp_path, report_type, date_from=day, date_to=day + timedelta(days=1))

        assert wb.sheetnames == ["Dashboard", "Notas", "Datos"]
        assert wb["Datos"].tables["datos"].ref == "A1:G4"  # encabezado + 3 pedidos del día



def make_dashboard_template(path, headers):
    """Plantilla mínima con lo que trae la real: tabla `datos` con una fila
    de ejemplo formateada, una hoja de dashboard con fórmula sobre la tabla
    y un gráfico, y una hoja extra que no debe tocarse."""
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"
    dashboard["A1"] = "Total vendido"
    dashboard["B1"] = "=SUM(datos[Valor])"
    notes = wb.create_sheet("Notas")
    notes["A1"] = "no tocar"

    data = wb.create_sheet("Datos")
    data.append(headers)
    data.append(["ejemplo"] * len(headers))
    for col, header in enumerate(headers, start=1):
        if header == "Valor":
            data.cell(row=2, column=col).number_format = '"$"#,##0'
        data.cell(row=3, column=col, value="basura vieja")
    data.add_table(Table(displayName="datos", ref=f"A1:{get_column_letter(len(headers))}3"))

    chart = BarChart()
    chart.add_data(Reference(data, min_col=1, min_row=1, max_row=3))
    dashboard.add_chart(chart, "D2")
    wb.save(path)
    return path


@pytest.fixture
def dashboard_template(tmp_path, monkeypatch):
    path = make_dashboard_template(
        tmp_path / "Plantilla_Dashboard.xlsx",
        ["Fecha", "Periodo", "Cliente", "Producto", "Cantidad", "Valor", "Estado"],
    )
    monkeypatch.setattr(dashboard_report, "DASHBOARD_TEMPLATE_PATH", path)
    return path


def test_dashboard_injects_orders_into_template_table(fake_store, tmp_path, dashboard_template):
    business = seed_report_data(fake_store)

    wb = generate(
        business,
        tmp_path,
        "dashboard",
        date_from=datetime(2026, 9, 1, tzinfo=DEFAULT_TZ),
        date_to=datetime(2026, 10, 1, tzinfo=DEFAULT_TZ),
    )

    ws = wb["Datos"]
    table = ws.tables["datos"]
    # 5 pedidos de septiembre (el de hace 40 días es de agosto): encabezado + 5.
    assert table.ref == "A1:G6"
    assert table.autoFilter.ref == "A1:G6"
    assert ws.max_row == 6  # las filas viejas de la plantilla no sobreviven
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2, max_row=6)]
    assert rows[0][0] == datetime(2026, 9, 8, 12, 30)  # ordenado por fecha, hora local sin zona
    assert {r[1] for r in rows} == {"2026-09"}
    assert {r[2] for r in rows} == {"Camilo", "Ana"}
    assert {r[3] for r in rows} == {"Menú del día", "Otras opciones"}
    assert sum(r[5] for r in rows) == 2 * 20500 + 3 * 18500 + 2 * 18500 + 20500 + 18500
    assert ws.cell(row=4, column=6).number_format == '"$"#,##0'  # formato de la plantilla

    # El resto de la plantilla queda intacto, y Excel recalcula al abrir.
    assert wb["Dashboard"]["B1"].value == "=SUM(datos[Valor])"
    assert len(wb["Dashboard"]._charts) == 1
    assert wb["Notas"]["A1"].value == "no tocar"
    assert wb.calculation.fullCalcOnLoad is True


def test_dashboard_never_includes_other_business_orders(fake_store, tmp_path, dashboard_template):
    business = seed_report_data(fake_store)
    other = seed_business(fake_store, name="Otro negocio")
    intruder = Orders()
    intruder.business_id = other.bussiness_id
    intruder.menu_id = "menu_del_dia"
    intruder.quantity = 99
    intruder.unit_price = 1
    intruder.total = 99
    intruder.status = "delivered"
    intruder.created_at = REFERENCE
    fake_store.seed(intruder)

    day = REFERENCE.replace(hour=0, minute=0)
    wb = generate(business, tmp_path, "dashboard", date_from=day, date_to=day + timedelta(days=1))

    quantities = [c.value for c in wb["Datos"]["E"][1:]]
    assert 99 not in quantities
    assert len(quantities) == 3


def test_dashboard_rejects_template_column_it_cannot_fill(fake_store, tmp_path, monkeypatch):
    business = seed_report_data(fake_store)
    path = make_dashboard_template(tmp_path / "t.xlsx", ["Fecha", "Categoria"])
    monkeypatch.setattr(dashboard_report, "DASHBOARD_TEMPLATE_PATH", path)

    with pytest.raises(ValueError, match="Categoria"):
        generate(
            business,
            tmp_path,
            "dashboard",
            date_from=REFERENCE - timedelta(days=60),
            date_to=REFERENCE + timedelta(days=1),
        )
